"""
Weekly name matcher: Odoo products still carrying a temporary ODOO-<id> code
in analytics_inventory_movements -> Wansoft products (owner, 2026-10-14).

Why a new matcher: the weekly mapping job (analysis/build_product_mapping.py)
only looks at Odoo products that are sold (sale_ok), so supplies such as
Lomo or Jitomate were never proposed, and its code-base rule matched Odoo
menu items (5100-/5200-) to Wansoft supplies (1000-), the 82 wrong mappings
reviewed on 2026-10-13.

Rule:
- Names are normalized: lower case, no accents, no punctuation, no Odoo code
  in brackets, no "(copia)", single spaces.
- Normalized names equal and exactly one Wansoft code with that name ->
  'approved' (mapping_source weekly_name_match:normalized_equal).
- Otherwise, Wansoft candidates sharing words (score = shared words / all
  words, at least 0.5) -> 'pending_review' with the best candidate, for the
  owner's review workbook; products with no candidate are only listed.
- Never touches an Odoo product that already has an 'approved' or 'rejected'
  dictionary row (owner decisions), and never approves a second code for one
  Odoo product (bug #44). An already approved row with the same Odoo name and
  Wansoft code but no Odoo id (old bridge-report approvals) only gets the id;
  a rejected one (rejected on 2026-10-13 for pointing at another Odoo id) is
  re-pointed to the right id and approved, with the history in notes.
- Wansoft catalog: dim_product (Wansoft codes) plus every product seen in
  getinputinventory_entrada, the live catalog.

The movements table picks the new codes up on its next Odoo rebuild
(product_code is a generated column).

Dry run by default (prints and writes the workbook); --apply writes the
dictionary.

    python -m scripts.match_odoo_products_by_name
    python -m scripts.match_odoo_products_by_name --apply
"""
import argparse
import os
import re
import unicodedata
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core.database.mysql import get_db_connection

REPORT_DIR = "reports/inventory_mapping_review"
MIN_SCORE = 0.5
SOURCE_EQUAL = "weekly_name_match:normalized_equal"
SOURCE_CANDIDATE = "weekly_name_match:candidate"


def normalize(name):
    text = unicodedata.normalize("NFKD", name or "").encode("ascii", "ignore").decode().lower()
    text = re.sub(r"\[[^\]]*\]", " ", text)           # Odoo code in brackets
    text = re.sub(r"\(copia\)", " ", text)
    text = re.sub(r"[^a-z0-9.]+", " ", text)           # keep decimals such as 3.78
    text = re.sub(r"(?<![0-9])\.|\.(?![0-9])", " ", text)
    return " ".join(text.split())


def words(normalized):
    return {w for w in normalized.split() if len(w) >= 3 or w.isdigit()}


def load_wansoft_catalog(cur):
    cur.execute("""
        SELECT wansoft_code, MAX(wansoft_product_name) FROM dim_product
        WHERE wansoft_code IS NOT NULL AND wansoft_code <> '' AND wansoft_product_name IS NOT NULL
        GROUP BY wansoft_code
        UNION
        SELECT CodigoProducto, MAX(NombreProducto) FROM getinputinventory_entrada
        WHERE CodigoProducto IS NOT NULL AND CodigoProducto <> '' GROUP BY CodigoProducto
    """)
    catalog = {}
    for code, name in cur.fetchall():
        catalog.setdefault(code, name)
    return catalog


def load_temporary_odoo_products(cur):
    cur.execute("""
        SELECT odoo_product_id, MAX(product_name), COUNT(*),
               SUM(final_source_status = 'odoo_final' AND counts_in_stock = 1)
        FROM analytics_inventory_movements
        WHERE source_system = 'odoo' AND product_code_status = 'temporal_odoo'
        GROUP BY odoo_product_id
    """)
    products = cur.fetchall()
    cur.execute("""
        SELECT DISTINCT odoo_product_id FROM inventory_mapping_dictionary
        WHERE odoo_product_id IS NOT NULL AND mapping_status IN ('approved', 'rejected')
    """)
    decided = {row[0] for row in cur.fetchall()}
    return [p for p in products if p[0] not in decided]


def match(products, catalog):
    by_name = {}
    indexed = []
    for code, name in catalog.items():
        norm = normalize(name)
        by_name.setdefault(norm, set()).add(code)
        indexed.append((code, name, words(norm)))
    approved, review, none = [], [], []
    for pid, name, moves, final_moves in products:
        norm = normalize(name)
        equal = by_name.get(norm, set())
        if len(equal) == 1:
            code = next(iter(equal))
            approved.append((pid, name, moves, final_moves, code, catalog[code], 1.0, ""))
            continue
        mine = words(norm)
        scored = sorted(
            ((len(mine & w) / len(mine | w), code, wname) for code, wname, w in indexed if mine and mine & w),
            reverse=True,
        )
        note = f"{len(equal)} códigos de Wansoft con el mismo nombre" if len(equal) > 1 else ""
        if scored and scored[0][0] >= MIN_SCORE:
            best = scored[0]
            others = "; ".join(f"{c} {n} ({s:.0%})" for s, c, n in scored[1:4])
            review.append((pid, name, moves, final_moves, best[1], best[2], round(best[0], 2), note or others))
        else:
            none.append((pid, name, moves, final_moves, "", "", 0, note))
    return approved, review, none


def write_workbook(approved, review, none):
    os.makedirs(REPORT_DIR, exist_ok=True)
    path = f"{REPORT_DIR}/emparejamiento_odoo_wansoft_{date.today():%Y%m%d}.xlsx"
    wb = Workbook()
    head = ["Odoo id", "Producto Odoo", "Movimientos", "Movimientos que cuentan", "Código Wansoft",
            "Producto Wansoft", "Coincidencia", "Otros candidatos / nota", "Decisión del dueño"]
    fill = PatternFill("solid", fgColor="035953")
    for title, rows in (("Aprobados automáticamente", approved), ("Para revisión", review), ("Sin candidato", none)):
        ws = wb.create_sheet(title)
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill
        for r in sorted(rows, key=lambda r: -(r[3] or 0)):
            ws.append(list(r) + [""])
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        for col, width in zip("ABCDEFGHI", (9, 40, 12, 14, 20, 40, 12, 60, 18)):
            ws.column_dimensions[col].width = width
    del wb["Sheet"]
    wb.save(path)
    return path


def save(conn, approved, review):
    cur = conn.cursor()
    sql = """
        INSERT INTO inventory_mapping_dictionary (
            domain, odoo_product_id, odoo_product_name, wansoft_code, wansoft_product_name,
            mapping_source, mapping_status, similarity_score, notes
        ) VALUES ('inventory', %s, %s, %s, %s, %s, %s, %s, %s)
        ON DUPLICATE KEY UPDATE
            -- An approved row with the same name and code but no Odoo id (old
            -- bridge-report approvals) only gets its id filled in.
            odoo_product_id = IF(mapping_status = 'rejected'
                                 OR (mapping_status = 'approved' AND odoo_product_id IS NOT NULL),
                                 odoo_product_id, VALUES(odoo_product_id)),
            mapping_source = IF(mapping_status IN ('approved', 'rejected'), mapping_source, VALUES(mapping_source)),
            similarity_score = IF(mapping_status IN ('approved', 'rejected'), similarity_score, VALUES(similarity_score)),
            notes = IF(mapping_status IN ('approved', 'rejected'), notes, VALUES(notes)),
            mapping_status = IF(mapping_status IN ('approved', 'rejected'), mapping_status, VALUES(mapping_status)),
            updated_at = IF(mapping_status IN ('approved', 'rejected'), updated_at, CURRENT_TIMESTAMP)
    """
    note = f"[{date.today()}] weekly name matcher"
    # A rejected row with the same Odoo name and Wansoft code was rejected for
    # its wrong Odoo id (2026-10-13), not for the name -> code pair: re-point
    # it to the right id and approve it, keeping the history in notes.
    repointed = 0
    for pid, name, _, _, code, _, _, _ in approved:
        cur.execute("""
            UPDATE inventory_mapping_dictionary
            SET notes = CONCAT_WS(' | ', NULLIF(notes, ''),
                    CONCAT('[', CURRENT_DATE, '] re-pointed from Odoo id ', COALESCE(odoo_product_id, 'none'),
                           ' (another product) to ', %s, ': normalized names equal')),
                odoo_product_id = %s, mapping_status = 'approved', mapping_source = %s,
                similarity_score = 100, updated_at = CURRENT_TIMESTAMP
            WHERE domain = 'inventory' AND odoo_product_name = %s AND wansoft_code = %s
              AND mapping_status = 'rejected'
        """, (pid, pid, SOURCE_EQUAL, name, code))
        repointed += cur.rowcount
    if repointed:
        print(f"Reapuntados (rechazados por id equivocado, nombre igual): {repointed}")
    rows = [(pid, name, code, wname, SOURCE_EQUAL, "approved", 100, note)
            for pid, name, _, _, code, wname, _, _ in approved]
    rows += [(pid, name, code, wname, SOURCE_CANDIDATE, "pending_review", round(score * 100, 2), note)
             for pid, name, _, _, code, wname, score, _ in review]
    cur.executemany(sql, rows)
    conn.commit()
    cur.close()
    return len(rows)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--apply", action="store_true", help="write the dictionary (default: dry run)")
    args = parser.parse_args(argv)

    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM information_schema.tables "
                "WHERE table_schema = DATABASE() AND table_name = 'analytics_inventory_movements'")
    if not cur.fetchone()[0]:
        print("[AVISO] analytics_inventory_movements no existe todavía: nada que emparejar.")
        conn.close()
        return 0
    catalog = load_wansoft_catalog(cur)
    products = load_temporary_odoo_products(cur)
    cur.close()
    approved, review, none = match(products, catalog)
    path = write_workbook(approved, review, none)
    total_final = sum(p[3] or 0 for p in products)
    covered = sum(r[3] or 0 for r in approved)
    print(f"Catálogo Wansoft: {len(catalog):,} códigos | productos Odoo con código temporal: {len(products):,}")
    print(f"  aprobados automáticamente: {len(approved):,} ({covered:,} de {total_final:,} movimientos que cuentan)")
    print(f"  para revisión: {len(review):,} | sin candidato: {len(none):,}")
    print(f"  libro de revisión: {path}")
    if args.apply:
        print(f"Diccionario: {save(conn, approved, review):,} filas escritas")
    else:
        print("Dry run: diccionario sin cambios. Usa --apply.")
    conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
